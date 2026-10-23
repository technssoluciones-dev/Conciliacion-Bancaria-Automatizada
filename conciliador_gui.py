import json
import os
import re
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

CONFIG_FILE = "mapeador_config.json"


def cargar_configuracion():
    if not os.path.exists(CONFIG_FILE):
        config_default = {
            "formatos_libro": {
                "Formato 5-1 Libro Diario (SUNAT / Estándar)": {
                    "skiprows": 5,
                    "col_fecha": "Fecha de la operación",
                    "col_id": "N°",
                    "col_glosa": "Glosa o Descripción de la operación",
                    "col_debe": "DEBE",
                    "col_haber": "HABER",
                    "col_cuenta": "Codigo",
                    "filtro_cuenta_prefijo": "104",
                    "tipo_monto": "DEBE_HABER",
                },
                "Estándar Genérico (Simple)": {
                    "skiprows": 0,
                    "col_fecha": "Fecha",
                    "col_id": "ID",
                    "col_glosa": "Descripción",
                    "col_debe": "Debe",
                    "col_haber": "Haber",
                    "col_cuenta": "Cuenta",
                    "tipo_monto": "DEBE_HABER",
                },
            }
        }
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(config_default, f, indent=4, ensure_ascii=False)
        return config_default
    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def extraer_metadatos_automaticos(ruta_banco, ruta_libro):
    meta = {
        "nombre": "N/A",
        "rut": "N/A",
        "banco": "N/A",
        "periodo": "N/A"
    }

    ext_banco = os.path.splitext(ruta_banco)[1].lower()
    fechas_encontradas = []

    if ext_banco == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(ruta_banco)
            texto_pdf_inicio = "\n".join([p.extract_text() for p in reader.pages[:2] if p.extract_text()])
            texto_completo = "\n".join([p.extract_text() for p in reader.pages if p.extract_text()])

            # Extraer RUT
            rut_match = re.search(r"\b(\d{1,2}\.?\d{3}\.?\d{3}-[\dkK])\b", texto_pdf_inicio)
            if rut_match:
                meta["rut"] = rut_match.group(1)

            # Detectar Banco
            texto_upper = texto_pdf_inicio.upper()
            if "SANTANDER" in texto_upper:
                meta["banco"] = "Banco Santander"
            elif "BCI" in texto_upper or "CREDITO E INVERSIONES" in texto_upper:
                meta["banco"] = "Banco BCI"
            elif "CHILE" in texto_upper and "BANCO" in texto_upper:
                meta["banco"] = "Banco de Chile"
            elif "ESTADO" in texto_upper:
                meta["banco"] = "BancoEstado"
            elif "SCOTIABANK" in texto_upper:
                meta["banco"] = "Scotiabank"
            elif "ITAU" in texto_upper or "ITAÚ" in texto_upper:
                meta["banco"] = "Banco Itaú"

            # Extraer Número de Cuenta
            cta_match = re.search(r"(?:CUENTA|CTA\.?|N°\s*DE\s*CUENTA)\s*(?:CORRIENTE|VISTA|N°|NO\.|#)?\s*([0-9\-]+)", texto_upper)
            if cta_match:
                cta_num = cta_match.group(1).strip()
                if len(cta_num) >= 4:
                    meta["periodo"] = f"Cta {cta_num}"

            # Extraer Nombre / Razón Social limpiando conectores redundantes
            for linea in texto_pdf_inicio.split("\n"):
                linea_u = linea.upper()
                if any(kw in linea_u for kw in ["SEÑOR(A)", "TITULAR", "CLIENTE", "RAZON SOCIAL", "RAZÓN SOCIAL", "NOMBRE"]):
                    # Se añade "O RAZON SOCIAL" y variantes al re.sub para eliminarlo por completo
                    clean = re.sub(
                        r"(SEÑOR\(A\)|TITULAR|CLIENTE|O\s+RAZON\s+SOCIAL|O\s+RAZÓN\s+SOCIAL|RAZON\s+SOCIAL|RAZÓN\s+SOCIAL|NOMBRE|RUC|RUT|:)", 
                        "", 
                        linea, 
                        flags=re.IGNORECASE
                    )
                    clean = re.sub(r"\b\d{1,2}\.?\d{3}\.?\d{3}-[\dkK]\b", "", clean).strip().upper()
                    if len(clean) > 3:
                        meta["nombre"] = clean
                        break

            fechas_encontradas = re.findall(r"\b(\d{2}/\d{2}/\d{4})\b", texto_completo)

        except Exception:
            pass

    # Analizar Libro Contable si no se halló la Razón Social en la cartola
    if meta["nombre"] == "N/A" or "RUC" in meta["nombre"] or "RUT" in meta["nombre"]:
        try:
            ext_libro = os.path.splitext(ruta_libro)[1].lower()
            df_lib_hdr = pd.read_excel(ruta_libro, nrows=15, header=None) if ext_libro != ".csv" else pd.read_csv(ruta_libro, nrows=15, header=None)
            for _, row in df_lib_hdr.iterrows():
                row_str = " ".join([str(val).strip() for val in row.values if pd.notna(val)])
                row_u = row_str.upper()
                if any(kw in row_u for kw in ["RAZON SOCIAL", "RAZÓN SOCIAL", "EMPRESA", "CLIENTE", "NOMBRE"]):
                    clean = re.sub(
                        r"(SEÑOR\(A\)|TITULAR|CLIENTE|O\s+RAZON\s+SOCIAL|O\s+RAZÓN\s+SOCIAL|RAZON\s+SOCIAL|RAZÓN\s+SOCIAL|EMPRESA|NOMBRE|RUT|RUC|:)", 
                        "", 
                        row_str, 
                        flags=re.IGNORECASE
                    )
                    clean = re.sub(r"\b\d{1,2}\.?\d{3}\.?\d{3}-[\dkK]\b", "", clean).strip().upper()
                    if len(clean) > 3:
                        meta["nombre"] = clean
                        break
        except Exception:
            pass

    # Rango de fechas para el campo Periodo
    if fechas_encontradas:
        fechas_dt = pd.to_datetime(fechas_encontradas, format="%d/%m/%Y", errors="coerce").dropna()
        if not fechas_dt.empty:
            f_min = fechas_dt.min().strftime("%d/%m/%Y")
            f_max = fechas_dt.max().strftime("%d/%m/%Y")
            str_rango = f"{f_min} al {f_max}" if f_min != f_max else f_min
            if meta["periodo"] != "N/A":
                meta["periodo"] += f" | {str_rango}"
            else:
                meta["periodo"] = str_rango

    return meta


def extraer_cartola_pdf(ruta_pdf):
    try:
        import pypdf

        reader = pypdf.PdfReader(ruta_pdf)
        full_text = "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
        lines = full_text.split("\n")

        movimientos = []
        date_pattern = re.compile(r"^(\d{2}/\d{2}/\d{4})\s+(.+)$")

        for line in lines:
            line = line.strip()
            match = date_pattern.match(line)
            if match:
                fecha = match.group(1)
                text = match.group(2)

                tokens = text.split()
                num_tokens = []
                text_tokens = []

                for t in reversed(tokens):
                    clean_t = t.replace(".", "")
                    if clean_t.isdigit():
                        num_tokens.insert(0, t)
                    else:
                        idx_split = len(tokens) - len(num_tokens)
                        text_tokens = tokens[:idx_split]
                        break

                descripcion = " ".join(text_tokens)

                if len(num_tokens) >= 2:
                    monto_str = num_tokens[0].replace(".", "").replace(",", ".")
                    monto = float(monto_str)
                    saldo_str = num_tokens[-1].replace(".", "").replace(",", ".")
                    saldo = float(saldo_str)

                    movimientos.append(
                        {
                            "Fecha": fecha,
                            "Descripción": descripcion,
                            "Monto_Raw": monto,
                            "Saldo": saldo,
                        }
                    )

        df = pd.DataFrame(movimientos)
        if df.empty:
            return pd.DataFrame()

        rows_final = []
        for i in range(len(df)):
            row = df.iloc[i]
            fecha = row["Fecha"]
            desc = row["Descripción"]
            monto = row["Monto_Raw"]
            saldo_actual = row["Saldo"]

            cargo = 0.0
            abono = 0.0

            if i > 0:
                saldo_pos = df.iloc[i - 1]["Saldo"]
                diff = saldo_pos - saldo_actual
                if abs(diff - (-monto)) < 1.0:
                    cargo = monto
                elif abs(diff - monto) < 1.0:
                    abono = monto
                else:
                    if any(
                        k in desc.lower()
                        for k in ["traspaso de:", "deposito", "abono", "tef recibida"]
                    ):
                        abono = monto
                    else:
                        cargo = monto
            else:
                cargo = monto

            rows_final.append(
                {
                    "Fecha": fecha,
                    "Descripción": desc,
                    "Cargos": cargo if cargo > 0 else 0.0,
                    "Abonos": abono if abono > 0 else 0.0,
                    "Monto_Neto": abono - cargo,
                }
            )

        return pd.DataFrame(rows_final)
    except Exception as e:
        raise Exception(f"Error al leer la cartola PDF: {str(e)}")


def cargar_archivo_banco(ruta):
    ext = os.path.splitext(ruta)[1].lower()

    if ext == ".pdf":
        return extraer_cartola_pdf(ruta)
    elif ext == ".csv":
        df = pd.read_csv(ruta)
    else:
        df = pd.read_excel(ruta)

    df.columns = [str(c).strip() for c in df.columns]

    col_cargo = next(
        (c for c in df.columns if any(k in c.lower() for k in ["cargo", "egreso", "débito"])),
        None,
    )
    col_abono = next(
        (c for c in df.columns if any(k in c.lower() for k in ["abono", "ingreso", "crédito"])),
        None,
    )
    col_monto = next(
        (c for c in df.columns if any(k in c.lower() for k in ["monto", "importe", "valor"])),
        None,
    )

    if col_cargo and col_abono:
        df[col_cargo] = pd.to_numeric(df[col_cargo], errors="coerce").fillna(0)
        df[col_abono] = pd.to_numeric(df[col_abono], errors="coerce").fillna(0)
        df["Monto_Neto"] = df[col_abono] - df[col_cargo]
    elif col_monto:
        df["Monto_Neto"] = pd.to_numeric(df[col_monto], errors="coerce").fillna(0)

    return df


def cargar_archivo_libro(ruta, config_libro):
    ext = os.path.splitext(ruta)[1].lower()

    if ext == ".csv":
        df_raw = pd.read_csv(ruta, header=None)
    else:
        df_raw = pd.read_excel(ruta, header=None)

    header_idx = None
    for idx, row in df_raw.iterrows():
        row_str = " ".join([str(v).upper() for v in row.values if pd.notna(v)])
        if "DEBE" in row_str and "HABER" in row_str:
            header_idx = idx
            break

    if header_idx is not None:
        if ext == ".csv":
            df = pd.read_csv(ruta, skiprows=header_idx)
        else:
            df = pd.read_excel(ruta, skiprows=header_idx)
    else:
        skip = config_libro.get("skiprows", 0)
        if ext == ".csv":
            df = pd.read_csv(ruta, skiprows=skip)
        else:
            df = pd.read_excel(ruta, skiprows=skip)

    df.columns = [str(c).strip() for c in df.columns]

    col_debe = next((c for c in df.columns if "debe" in c.lower()), None)
    col_haber = next((c for c in df.columns if "haber" in c.lower()), None)
    col_cuenta = next(
        (c for c in df.columns if any(k in c.lower() for k in ["codigo", "código", "cuenta"])),
        None,
    )
    prefijo = config_libro.get("filtro_cuenta_prefijo", "104")

    if col_cuenta and col_cuenta in df.columns:
        df = df[df[col_cuenta].astype(str).str.startswith(str(prefijo))]

    if col_debe and col_haber:
        df[col_debe] = pd.to_numeric(df[col_debe], errors="coerce").fillna(0)
        df[col_haber] = pd.to_numeric(df[col_haber], errors="coerce").fillna(0)
        df["Monto_Neto"] = df[col_debe] - df[col_haber]
    else:
        num_cols = df.select_dtypes(include=["number"]).columns
        if len(num_cols) >= 2:
            df["Monto_Neto"] = df[num_cols[0]] - df[num_cols[1]]

    return df


def procesar_conciliacion(ruta_banco, ruta_libro, tipo_formato_libro, ruta_salida):
    config = cargar_configuracion()
    cfg_libro = config["formatos_libro"].get(tipo_formato_libro, {})

    info_empresa = extraer_metadatos_automaticos(ruta_banco, ruta_libro)

    df_banco = cargar_archivo_banco(ruta_banco)
    df_libro = cargar_archivo_libro(ruta_libro, cfg_libro)

    if "Monto_Neto" not in df_banco.columns or "Monto_Neto" not in df_libro.columns:
        raise ValueError("No se pudieron determinar las columnas numéricas para el cruce.")

    df_cruce = pd.merge(
        df_banco,
        df_libro,
        on="Monto_Neto",
        how="outer",
        suffixes=("_Banco", "_Libro"),
        indicator=True,
    )

    conciliados = df_cruce[df_cruce["_merge"] == "both"].copy()
    solo_banco = df_cruce[df_cruce["_merge"] == "left_only"].copy()
    solo_libro = df_cruce[df_cruce["_merge"] == "right_only"].copy()

    for df_temp in [conciliados, solo_banco, solo_libro]:
        cols_to_drop = [c for c in df_temp.columns if "Unnamed" in str(c) or c == "_merge"]
        df_temp.drop(columns=cols_to_drop, errors="ignore", inplace=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Informe de Conciliación"
    ws.views.sheetView[0].showGridLines = True

    font_title_main = Font(name="Calibri", size=16, bold=True, color="1F4E78")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="595959")
    font_sec_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_col_header = Font(name="Calibri", size=10, bold=True, color="1F4E78")
    font_data = Font(name="Calibri", size=10)
    font_data_bold = Font(name="Calibri", size=10, bold=True)

    fill_navy = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_sub_hdr = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FBFD", end_color="F9FBFD", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    border_top_double = Border(
        top=Side(style="thin", color="000000"),
        bottom=Side(style="double", color="000000"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    ws.merge_cells("A1:E1")
    ws["A1"] = "INFORME DE CONCILIACIÓN BANCARIA"
    ws["A1"].font = font_title_main
    ws["A1"].alignment = align_left

    ws["A3"] = "Razón Social / Nombre:"
    ws["B3"] = info_empresa.get("nombre", "N/A")
    ws["A4"] = "RUT / ID Fiscal:"
    ws["B4"] = info_empresa.get("rut", "N/A")
    ws["D3"] = "Institución Bancaria:"
    ws["E3"] = info_empresa.get("banco", "N/A")
    ws["D4"] = "N° Cuenta / Periodo:"
    ws["E4"] = info_empresa.get("periodo", "N/A")

    for r in [3, 4]:
        ws[f"A{r}"].font = font_data_bold
        ws[f"D{r}"].font = font_data_bold
        ws[f"B{r}"].font = font_data
        ws[f"E{r}"].font = font_data

    current_row = 6

    def escribir_seccion(titulo, df_data, row_start):
        ws.merge_cells(start_row=row_start, start_column=1, end_row=row_start, end_column=5)
        cell_sec = ws.cell(row=row_start, column=1, value=f"  {titulo.upper()} ({len(df_data)} registros)")
        cell_sec.font = font_sec_header
        cell_sec.fill = fill_navy
        cell_sec.alignment = align_left
        ws.row_dimensions[row_start].height = 24
        
        row_start += 1

        headers = ["Fecha", "Descripción / Glosa", "Cargos / Débitos", "Abonos / Créditos", "Monto Neto"]
        for c_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=row_start, column=c_idx, value=h_text)
            cell.font = font_col_header
            cell.fill = fill_sub_hdr
            cell.alignment = align_center
            cell.border = border_thin
        
        ws.row_dimensions[row_start].height = 20
        row_start += 1

        if df_data.empty:
            ws.merge_cells(start_row=row_start, start_column=1, end_row=row_start, end_column=5)
            c_empty = ws.cell(row=row_start, column=1, value="No se encontraron registros para esta sección.")
            c_empty.font = font_subtitle
            c_empty.alignment = align_center
            row_start += 1
            return row_start + 1

        start_data_row = row_start
        for idx, (_, row) in enumerate(df_data.iterrows()):
            ws.row_dimensions[row_start].height = 19
            
            fecha_val = row.get("Fecha_Banco", row.get("Fecha_Libro", row.get("Fecha", "")))
            desc_val = row.get("Descripción_Banco", row.get("Descripción_Libro", row.get("Descripción", row.get("Glosa o Descripción de la operación", ""))))
            cargos_val = row.get("Cargos", row.get("DEBE", 0.0))
            abonos_val = row.get("Abonos", row.get("HABER", 0.0))
            monto_val = row.get("Monto_Neto", 0.0)

            c1 = ws.cell(row=row_start, column=1, value=fecha_val)
            c2 = ws.cell(row=row_start, column=2, value=str(desc_val) if pd.notna(desc_val) else "")
            c3 = ws.cell(row=row_start, column=3, value=float(cargos_val) if pd.notna(cargos_val) else 0.0)
            c4 = ws.cell(row=row_start, column=4, value=float(abonos_val) if pd.notna(abonos_val) else 0.0)
            c5 = ws.cell(row=row_start, column=5, value=float(monto_val) if pd.notna(monto_val) else 0.0)

            c1.alignment = align_center
            c2.alignment = align_left
            c3.alignment = align_right
            c4.alignment = align_right
            c5.alignment = align_right

            c3.number_format = "#,##0.00"
            c4.number_format = "#,##0.00"
            c5.number_format = "#,##0.00"

            for cell in [c1, c2, c3, c4, c5]:
                cell.font = font_data
                cell.border = border_thin
                if idx % 2 == 1:
                    cell.fill = fill_zebra

            row_start += 1

        end_data_row = row_start - 1
        
        ws.cell(row=row_start, column=2, value="TOTAL SECCIÓN:").font = font_data_bold
        ws.cell(row=row_start, column=2).alignment = align_right
        
        t_cargos = ws.cell(row=row_start, column=3, value=f"=SUM(C{start_data_row}:C{end_data_row})")
        t_abonos = ws.cell(row=row_start, column=4, value=f"=SUM(D{start_data_row}:D{end_data_row})")
        t_neto = ws.cell(row=row_start, column=5, value=f"=SUM(E{start_data_row}:E{end_data_row})")

        for c_tot in [t_cargos, t_abonos, t_neto]:
            c_tot.font = font_data_bold
            c_tot.number_format = "#,##0.00"
            c_tot.alignment = align_right
            c_tot.border = border_top_double

        ws.row_dimensions[row_start].height = 22
        return row_start + 3

    current_row = escribir_seccion("1. Movimientos Conciliados (Coinciden en Banco y Libro)", conciliados, current_row)
    current_row = escribir_seccion("2. Pendientes en Banco (Solo figuran en Extracto Bancario)", solo_banco, current_row)
    current_row = escribir_seccion("3. Pendientes en Libro (Solo figuran en Libro Contable)", solo_libro, current_row)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18

    wb.save(ruta_salida)


class ConciliadorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Sistema Automatizado de Conciliación Bancaria")
        self.root.geometry("580x280")
        self.root.resizable(False, False)

        self.config = cargar_configuracion()
        self.create_widgets()

    def create_widgets(self):
        frame = ttk.LabelFrame(self.root, text=" Archivos de Entrada ", padding=(15, 15))
        frame.pack(fill="both", expand=True, padx=15, pady=15)

        ttk.Label(frame, text="Formato Libro Contable:").grid(row=0, column=0, sticky="w", pady=5)
        self.combo_formato = ttk.Combobox(
            frame,
            values=list(self.config["formatos_libro"].keys()),
            state="readonly",
            width=36,
        )
        self.combo_formato.current(0)
        self.combo_formato.grid(row=0, column=1, columnspan=2, sticky="w", pady=5)

        ttk.Label(frame, text="Libro Contable:").grid(row=1, column=0, sticky="w", pady=8)
        self.entry_libro = ttk.Entry(frame, width=36)
        self.entry_libro.grid(row=1, column=1, pady=8)
        ttk.Button(frame, text="Examinar...", command=self.seleccionar_libro).grid(row=1, column=2, padx=5, pady=8)

        ttk.Label(frame, text="Cartola Bancaria:").grid(row=2, column=0, sticky="w", pady=8)
        self.entry_banco = ttk.Entry(frame, width=36)
        self.entry_banco.grid(row=2, column=1, pady=8)
        ttk.Button(frame, text="Examinar...", command=self.seleccionar_banco).grid(row=2, column=2, padx=5, pady=8)

        btn_ejecutar = ttk.Button(self.root, text="⚡ PROCESAR Y GENERAR REPORTE AUTOMÁTICO", command=self.ejecutar)
        btn_ejecutar.pack(fill="x", padx=15, pady=(0, 15))

    def seleccionar_libro(self):
        ftypes = [("Todos los soportados", "*.xlsx *.xls *.csv *.pdf"), ("Excel", "*.xlsx *.xls"), ("CSV", "*.csv"), ("PDF", "*.pdf")]
        filename = filedialog.askopenfilename(title="Seleccionar Libro Contable", filetypes=ftypes)
        if filename:
            self.entry_libro.delete(0, tk.END)
            self.entry_libro.insert(0, filename)

    def seleccionar_banco(self):
        ftypes = [("Todos los soportados", "*.xlsx *.xls *.csv *.pdf"), ("PDF", "*.pdf"), ("Excel", "*.xlsx *.xls"), ("CSV", "*.csv")]
        filename = filedialog.askopenfilename(title="Seleccionar Cartola Bancaria", filetypes=ftypes)
        if filename:
            self.entry_banco.delete(0, tk.END)
            self.entry_banco.insert(0, filename)

    def ejecutar(self):
        ruta_libro = self.entry_libro.get()
        ruta_banco = self.entry_banco.get()
        formato = self.combo_formato.get()

        if not ruta_libro or not ruta_banco:
            messagebox.showwarning("Atención", "Por favor selecciona ambos archivos antes de continuar.")
            return

        ruta_salida = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")],
            title="Guardar Conciliación Automatizada",
        )

        if not ruta_salida:
            return

        try:
            procesar_conciliacion(ruta_banco, ruta_libro, formato, ruta_salida)
            messagebox.showinfo(
                "Éxito",
                f"✅ Análisis automático completado.\n\nEl archivo final fue guardado en:\n{ruta_salida}",
            )
        except Exception as e:
            messagebox.showerror("Error", f"Ocurrió un error en el proceso:\n{str(e)}")


if __name__ == "__main__":
    root = tk.Tk()
    app = ConciliadorApp(root)
    root.mainloop()